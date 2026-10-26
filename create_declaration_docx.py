from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path('/Users/ksyeo/Desktop/tp/NST2060_2610_Declaration.docx')


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = 'Arial'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        if edge in kwargs:
            tag = 'w:{}'.format(edge)
            element = borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                borders.append(element)
            for key, value in kwargs[edge].items():
                element.set(qn('w:{}'.format(key)), str(value))


def add_text(paragraph, text, **style):
    run = paragraph.add_run(text)
    set_run_font(run, **style)
    return run


def add_body_paragraph(doc, text='', space_after=8):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = 1.15
    add_text(paragraph, text, size=12)
    return paragraph


def add_bullet(doc, parts):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.85)
    paragraph.paragraph_format.first_line_indent = Cm(-0.5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.12
    add_text(paragraph, '🧬  ', size=12)
    for text, style in parts:
        add_text(paragraph, text, size=12, **style)
    return paragraph


doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(1.45)
section.bottom_margin = Cm(1.35)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Arial')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
normal.font.size = Pt(12)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(0)
add_text(title, 'NST2060 • Knowing Me, Knowing You: Venturing into the Genomic Age', size=15, bold=True)

heading = doc.add_paragraph()
heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
heading.paragraph_format.space_after = Pt(22)
add_text(heading, '🧬  ', size=16)
add_text(heading, 'Declaration of Academic Integrity', size=20, bold=True)
add_text(heading, '  🧬', size=16)

add_body_paragraph(
    doc,
    'Please append this declaration to both of your assignment submissions (individual), and also\n'
    'to the final group project report submission (one copy for the group).',
    space_after=15,
)
add_body_paragraph(doc, 'By signing this document, I/we:', space_after=12)

add_bullet(doc, [
    ('Declare that the submitted work listed below has been completed in accordance with\n'
     'all the relevant academic integrity rules of NUS, with particular attention to the ', {}),
    ('NUS\nPlagiarism Policy', {'color': (0, 112, 192)}),
])
add_bullet(doc, [
    ('Recognize that violation of said rules will result in being subject to the relevant\n'
     'disciplinary proceedings of NUS', {}),
])
add_bullet(doc, [
    ('Understand that in cases of suspected ', {}),
    ('unacknowledged', {'italic': True}),
    (' AI use, the instructor reserves\n'
     'the right to ask students to orally explain their work, in order to verify and affirm their\n'
     'comprehension of their written submission', {}),
])

def label(doc, label_text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(11)
    paragraph.paragraph_format.space_after = Pt(5)
    add_text(paragraph, label_text, size=12, bold=True, italic=True)


def fill_line(doc, main_text, hint):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16)
    cell.vertical_alignment = 1
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.left_indent = Cm(0.2)
    add_text(paragraph, main_text, size=12)
    add_text(paragraph, '  ', size=12)
    add_text(paragraph, hint, size=12, italic=True, color=(128, 128, 128))
    set_cell_border(cell, bottom={'val': 'single', 'sz': '8', 'color': '000000'})


label(doc, 'Submission type:')
fill_line(doc, 'Assignment 1 / Assignment 2 / Group project report', '(delete accordingly)')

label(doc, 'Use of generative AI:')
fill_line(doc, 'Transcript attached / Declaring no generative AI used', '(delete accordingly)')

label(doc, 'Date:')
date_table = doc.add_table(rows=1, cols=1)
date_table.autofit = False
date_cell = date_table.cell(0, 0)
date_cell.width = Cm(4.4)
date_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
set_cell_border(date_cell, bottom={'val': 'single', 'sz': '8', 'color': '000000'})

label(doc, 'Names:')
for student in range(1, 5):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    label_cell, line_cell = table.rows[0].cells
    label_cell.width = Cm(2.2)
    line_cell.width = Cm(8.8)
    paragraph = label_cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    add_text(paragraph, f'Student {student}:', size=12)
    line_cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    set_cell_border(line_cell, bottom={'val': 'single', 'sz': '8', 'color': '000000'})

# Footer
footer = section.footer
footer_paragraph = footer.paragraphs[0]
footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_text(footer_paragraph, '🧬  Page 1 of 1', size=10, color=(128, 128, 128))

doc.save(OUTPUT)
print(OUTPUT)
