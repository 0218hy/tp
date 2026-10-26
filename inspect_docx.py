from docx import Document

doc = Document('/Users/ksyeo/Desktop/tp/NST2060 Assignment 2 formatted.docx')

for i, paragraph in enumerate(doc.paragraphs):
    print(f'P{i}: {paragraph.text!r}')
    for j, run in enumerate(paragraph.runs):
        print(
            f'  R{j}: {run.text!r} '
            f'sub={run.font.subscript} sup={run.font.superscript}'
        )

for table_index, table in enumerate(doc.tables):
    print(f'TABLE {table_index}')
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            print(f' C{row_index},{cell_index}: {cell.text!r}')
