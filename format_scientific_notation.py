from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


SOURCE = Path('/Users/ksyeo/Downloads/NST2060 Assignment 2.docx')
OUTPUT = Path('/Users/ksyeo/Desktop/tp/NST2060 Assignment 2 formatted.docx')


# Ordered from most-specific to least-specific so, for example, r2AB is
# handled as r-squared with an AB subscript rather than as r-squared plus text.
TOKEN_RE = re.compile(
    r'(D(?:AB|AC|BC)2|r2(?:AB|AC|BC)?|F(?:ST|st)|Hobs|D(?:AB|AC|BC)|'
    r'[pqh]2|D2|H(?:T|S|A|s)\b|V(?:A|E)\b)'
)


def set_text(run_element, text):
    """Replace text-bearing children while preserving the source run properties."""
    for child in list(run_element):
        if child.tag != qn('w:rPr'):
            run_element.remove(child)
    text_element = OxmlElement('w:t')
    if text.startswith(' ') or text.endswith(' '):
        text_element.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    text_element.text = text
    run_element.append(text_element)


def set_vertical_alignment(run_element, alignment):
    properties = run_element.find(qn('w:rPr'))
    if properties is None:
        properties = OxmlElement('w:rPr')
        run_element.insert(0, properties)
    for old in list(properties.findall(qn('w:vertAlign'))):
        properties.remove(old)
    vertical_align = OxmlElement('w:vertAlign')
    vertical_align.set(qn('w:val'), alignment)
    properties.append(vertical_align)


def pieces_for_token(token):
    """Return (text, vertical-alignment) pieces for a scientific token."""
    if token.startswith('r2'):
        suffix = token[2:]
        pieces = [('r', None), ('2', 'superscript')]
        if suffix:
            pieces.append((suffix, 'subscript'))
        return pieces
    if token.startswith('D') and token.endswith('2') and token[1:-1] in {'AB', 'AC', 'BC'}:
        return [('D', None), (token[1:-1], 'subscript'), ('2', 'superscript')]
    if token.startswith('D') and token[1:] in {'AB', 'AC', 'BC'}:
        return [('D', None), (token[1:], 'subscript')]
    if token == 'D2':
        return [('D', None), ('2', 'superscript')]
    if token in {'p2', 'q2', 'h2'}:
        return [(token[0], None), ('2', 'superscript')]
    if token in {'FST', 'Fst'}:
        return [('F', None), (token[1:], 'subscript')]
    if token == 'Hobs':
        return [('H', None), ('obs', 'subscript')]
    if token.startswith('H'):
        return [('H', None), (token[1:], 'subscript')]
    if token in {'VA', 'VE'}:
        return [('V', None), (token[1:], 'subscript')]
    raise ValueError(f'Unhandled token: {token}')


def replace_run(run):
    text = run.text
    if not text or not TOKEN_RE.search(text):
        return 0

    parent = run._r.getparent()
    insertion_index = parent.index(run._r)
    fragments = []
    cursor = 0
    changes = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > cursor:
            fragments.append((text[cursor:match.start()], None))
        fragments.extend(pieces_for_token(match.group(0)))
        cursor = match.end()
        changes += 1
    if cursor < len(text):
        fragments.append((text[cursor:], None))

    for fragment_text, alignment in fragments:
        if not fragment_text:
            continue
        new_run = deepcopy(run._r)
        set_text(new_run, fragment_text)
        if alignment:
            set_vertical_alignment(new_run, alignment)
        parent.insert(insertion_index, new_run)
        insertion_index += 1
    parent.remove(run._r)
    return changes


def iter_paragraphs(container):
    for paragraph in container.paragraphs:
        yield paragraph
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from iter_paragraphs(cell)


document = Document(SOURCE)
changes = 0
for paragraph in iter_paragraphs(document):
    for run in list(paragraph.runs):
        changes += replace_run(run)

for section in document.sections:
    for header_or_footer in (section.header, section.footer):
        for paragraph in iter_paragraphs(header_or_footer):
            for run in list(paragraph.runs):
                changes += replace_run(run)

document.save(OUTPUT)
print(f'Formatted {changes} scientific-notation tokens.')
print(OUTPUT)
